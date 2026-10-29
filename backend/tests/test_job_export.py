from datetime import datetime
from copy import copy
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.db import SessionLocal
from app.imports import CAMPUS_HEADERS, HEADERS
from app.models import Job, ImportBatch
from test_campus_import import campus_file, record
from test_business import excel


def styled_source(records):
    wb = load_workbook(BytesIO(campus_file(records, list(reversed(CAMPUS_HEADERS)))))
    sheet = wb.worksheets[0]
    sheet.column_dimensions['A'].width = 37
    sheet.column_dimensions['C'].width = 45
    sheet.row_dimensions[1].height = 31
    sheet.row_dimensions[2].height = 47
    sheet.row_dimensions[3].height = 63
    sheet.freeze_panes = 'B2'
    sheet.auto_filter.ref = f'A1:T{len(records)+1}'
    sheet['A1'].font = Font(name='微软雅黑', bold=True, color='AABBCC')
    sheet['A1'].fill = PatternFill('solid', fgColor='123456')
    sheet['A1'].alignment = Alignment(wrap_text=True)
    date_column = list(reversed(CAMPUS_HEADERS)).index('更新时间') + 1
    sheet.cell(3, date_column, datetime(2026, 10, 26)).number_format = 'yyyy/mm/dd'
    link_column = list(reversed(CAMPUS_HEADERS)).index('简历投递链接') + 1
    sheet.cell(3, link_column).hyperlink = 'https://example.com/original'
    output = BytesIO(); wb.save(output)
    return output.getvalue()


def import_source(client, content):
    response = client.post('/api/imports/preview', files={'file': ('岗位来源.xlsx', content)})
    assert response.status_code == 201, response.text
    batch = response.json()
    assert 'source_file' not in batch and 'source_sheet' not in batch
    assert client.post(f"/api/imports/{batch['id']}/commit").status_code == 200
    return batch


def test_export_preserves_original_selected_row_styles_and_types(clients):
    c = clients('alice')
    content = styled_source([record(**{'公司': '不能导出的公司'}), record(**{'公司': '应导出公司'})])
    batch = import_source(c, content)
    jobs = c.get('/api/jobs').json()['items']
    selected = next(job for job in jobs if job['company'] == '应导出公司')
    assert 'source_batch_id' not in selected
    sid = c.post('/api/students', json={'name': '张三'}).json()['id']
    response = c.post(f'/api/students/{sid}/jobs/export', json={'job_ids': [selected['id'], selected['id']]})
    assert response.status_code == 200, response.text
    assert "filename*=UTF-8''" in response.headers['content-disposition']
    wb = load_workbook(BytesIO(response.content))
    assert len(wb.worksheets) == 1
    sheet = wb.active
    assert sheet.title == '校招汇总表（优先）'
    assert sheet.max_row == 2 and sheet.max_column == 20
    headers = list(next(sheet.values))
    assert headers == list(reversed(CAMPUS_HEADERS))
    assert sheet.cell(2, headers.index('公司') + 1).value == '应导出公司'
    assert '不能导出的公司' not in str(list(sheet.values))
    original = load_workbook(BytesIO(content)).worksheets[0]
    assert copy(sheet['A1'].font) == copy(original['A1'].font)
    assert copy(sheet['A1'].fill) == copy(original['A1'].fill)
    assert copy(sheet['A1'].alignment) == copy(original['A1'].alignment)
    assert sheet.column_dimensions['A'].width == 37
    assert sheet.row_dimensions[2].height == 63
    assert sheet.freeze_panes == 'B2'
    assert sheet.auto_filter.ref == 'A1:T2'
    cell = sheet.cell(2, headers.index('更新时间') + 1)
    assert cell.value == datetime(2026, 10, 26) and cell.number_format == 'yyyy/mm/dd'
    assert sheet.cell(2, headers.index('简历投递链接') + 1).hyperlink.target == 'https://example.com/original'
    assert c.get(f'/api/students/{sid}').json()['recommendations'] == []
    with SessionLocal() as db:
        source = db.get(ImportBatch, batch['id'])
        assert source.source_file == content
        job = db.get(Job, selected['id'])
        assert job.source_batch_id == batch['id'] and job.source_row_number == 3


def test_export_groups_different_files_and_supports_cross_consultant_jobs(clients):
    alice, bob = clients('alice'), clients('bob')
    import_source(alice, campus_file([record(**{'公司': '甲公司'})]))
    import_source(bob, campus_file([record(**{'公司': '乙公司'})]))
    jobs = alice.get('/api/jobs').json()['items']
    sid = alice.post('/api/students', json={'name': '跨来源学员'}).json()['id']
    response = alice.post(f'/api/students/{sid}/jobs/export', json={'job_ids': [j['id'] for j in jobs]})
    assert response.status_code == 200, response.text
    wb = load_workbook(BytesIO(response.content))
    assert len(wb.worksheets) == 2 and len(set(wb.sheetnames)) == 2
    assert {sheet['B2'].value for sheet in wb} == {'甲公司', '乙公司'}
    assert all(sheet.max_row == 2 for sheet in wb)


def test_export_owner_permissions_validation_and_legacy_fallback(clients):
    alice, bob, admin = clients('alice'), clients('bob'), clients('admin')
    import_source(alice, excel([['后端工程师', '历史公司', '上海']]))
    job = alice.get('/api/jobs').json()['items'][0]
    # Simulate a record imported before source retention existed.
    with SessionLocal() as db:
        saved = db.get(Job, job['id']); saved.source_batch_id = None; saved.source_row_number = None
        db.commit()
    sid = alice.post('/api/students', json={'name': '导出学员'}).json()['id']
    url = f'/api/students/{sid}/jobs/export'
    body = {'job_ids': [job['id']]}
    assert bob.post(url, json=body).status_code == 404
    assert admin.post(url, json=body).status_code == 200
    assert alice.post(url, json={'job_ids': []}).status_code == 422
    assert alice.post(url, json={'job_ids': [999999]}).status_code == 404
    response = alice.post(url, json=body)
    sheet = load_workbook(BytesIO(response.content)).active
    assert list(next(sheet.values)) == list(HEADERS)
    assert sheet['A2'].value == '后端工程师'


def test_historical_campus_fields_and_formula_text_are_preserved_safely(clients):
    c = clients('alice')
    import_source(c, campus_file([record()]))
    job = c.get('/api/jobs').json()['items'][0]
    with SessionLocal() as db:
        saved = db.get(Job, job['id'])
        saved.source_batch_id = None; saved.source_row_number = None
        saved.recruitment_details = {**saved.recruitment_details, '备注': '=HYPERLINK("bad")'}
        db.commit()
    sid = c.post('/api/students', json={'name': '历史字段'}).json()['id']
    response = c.post(f'/api/students/{sid}/jobs/export', json={'job_ids': [job['id']]})
    sheet = load_workbook(BytesIO(response.content)).active
    assert list(next(sheet.values)) == CAMPUS_HEADERS
    assert sheet['J2'].value == '招满为止'
    assert sheet['O2'].value.startswith("'=") and sheet['O2'].data_type == 's'
    assert sheet['K2'].hyperlink.target.startswith('https://example.com/apply')
