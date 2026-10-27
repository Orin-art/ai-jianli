from io import BytesIO
from datetime import timedelta
from concurrent.futures import ThreadPoolExecutor
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from app.main import app, today
from app.imports import HEADERS
from app.db import SessionLocal
from app.models import ImportBatch, AuditLog
from sqlalchemy import select

def excel(rows, headers=None):
    wb = Workbook(); ws = wb.active; ws.append(headers or list(HEADERS))
    for row in rows:
        ws.append(row)
    data = BytesIO(); wb.save(data); return data.getvalue()

def preview(client, rows):
    return client.post('/api/imports/preview', files={'file': ('岗位.xlsx', excel(rows))})

def add_job(client, title='后端工程师'):
    batch = preview(client, [[title, '测试科技', '上海', '本科', '应届生', '计算机', 8000, 12000, '开发接口', '测试来源']]).json()
    assert client.post(f"/api/imports/{batch['id']}/commit").status_code == 200
    return client.get('/api/jobs').json()['items'][0]

def test_auth_csrf_and_revocation(clients):
    outsider = TestClient(app)
    assert outsider.get('/api/jobs').status_code == 401
    assert outsider.post('/api/auth/login', json={'username': 'admin', 'password': 'anything'}).status_code == 403
    alice = clients('alice'); admin = clients('admin')
    assert 'HttpOnly' in alice.cookies.get('career_session', '') or alice.cookies.get('career_session')
    assert alice.get('/api/users').status_code == 403
    assert alice.post('/api/users', json={'username': 'new', 'name': '新顾问', 'password': 'Some-password-123'}).status_code == 403
    assert admin.patch('/api/users/2', json={'active': False}).status_code == 200
    assert alice.get('/api/auth/me').status_code == 401
    assert admin.post('/api/auth/logout').status_code == 200
    assert admin.get('/api/auth/me').status_code == 401

def test_login_lockout_and_redacted_validation():
    client = TestClient(app, headers={'X-Requested-With': 'career-web'})
    for _ in range(5):
        assert client.post('/api/auth/login', json={'username': 'admin', 'password': 'wrong'}).status_code == 401
    assert client.post('/api/auth/login', json={'username': 'admin', 'password': 'wrong'}).status_code == 429
    response = client.post('/api/auth/login', json={'username': 'admin', 'password': 's' * 129})
    assert response.status_code == 422 and 's' * 129 not in response.text

def test_import_validation_duplicates_and_commit(clients):
    c = clients('alice')
    rows = [['后端开发', '甲公司', '上海', '本科', '应届生', '计算机', 8000, 10000], ['后端开发', '甲公司', '上海'], ['', '乙公司', '北京'], ['前端开发', '丙公司', '深圳', '本科', '不限', '', 12000, 8000], ['=1+1', '公式公司', '上海']]
    response = preview(c, rows)
    assert response.status_code == 201, response.text
    b = response.json(); assert b['counts'] == {'可导入': 1, '错误': 3, '重复': 1, '已导入': 0}
    assert c.get('/api/jobs').json()['total'] == 0
    errors = c.get(f"/api/imports/{b['id']}/errors")
    sheet = load_workbook(BytesIO(errors.content)).active
    assert sheet.max_row == 5
    assert all(cell.data_type != 'f' for row in sheet for cell in row)
    committed = c.post(f"/api/imports/{b['id']}/commit")
    assert committed.status_code == 200, committed.text
    assert committed.json()['counts']['已导入'] == 1
    assert c.post(f"/api/imports/{b['id']}/commit").status_code == 409
    assert c.get('/api/jobs').json()['total'] == 1
    assert preview(c, [rows[0]]).json()['counts']['重复'] == 1

def test_import_race_and_preview_ownership(clients):
    alice = clients('alice'); bob = clients('bob')
    a = preview(alice, [['开发', '甲', '上海']]).json()
    b = preview(bob, [['开发', '甲', '上海']]).json()
    assert bob.get(f"/api/imports/{a['id']}").status_code == 404
    assert bob.get(f"/api/imports/{a['id']}/errors").status_code == 404
    assert bob.post(f"/api/imports/{a['id']}/commit").status_code == 404
    assert alice.post(f"/api/imports/{a['id']}/commit").status_code == 200
    result = bob.post(f"/api/imports/{b['id']}/commit").json()
    assert result['counts']['重复'] == 1 and result['counts']['已导入'] == 0

def test_reject_bad_workbooks(clients):
    c = clients('alice')
    assert c.post('/api/imports/preview', files={'file': ('岗位.xls', b'not excel')}).status_code == 400
    assert c.post('/api/imports/preview', files={'file': ('岗位.xlsx', b'not excel')}).status_code == 400
    assert c.post('/api/imports/preview', files={'file': ('岗位.xlsx', excel([], ['未知列']))}).status_code == 400
    assert preview(c, []).status_code == 400
    assert preview(c, [['岗位', '企业', '城市']] * 2001).status_code == 400
    assert c.post('/api/imports/preview', content=b'x' * (11 * 1024 * 1024 + 1)).status_code == 413

def test_concurrent_confirm_has_one_winner(clients):
    a, b = clients('alice'), clients('alice')
    batch = preview(a, [['并发测试', '测试公司', '上海']]).json()
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda c: c.post(f"/api/imports/{batch['id']}/commit").status_code, [a, b]))
    assert sorted(results) == [200, 409]
    assert a.get('/api/jobs').json()['total'] == 1

def test_cancel_import_including_duplicate_batches(clients):
    a, b, admin = clients('alice'), clients('bob'), clients('admin')
    original = preview(a, [['重复岗位', '测试企业', '上海']]).json()
    assert a.post(f"/api/imports/{original['id']}/commit").status_code == 200
    assert a.post(f"/api/imports/{original['id']}/cancel").status_code == 409
    duplicate = preview(a, [['重复岗位', '测试企业', '上海']]).json()
    assert duplicate['counts']['可导入'] == 0 and duplicate['counts']['重复'] == 1
    url = f"/api/imports/{duplicate['id']}"
    assert b.post(url + '/cancel').status_code == 404
    response = a.post(url + '/cancel')
    assert response.status_code == 200 and response.json()['status'] == '已取消'
    assert response.json()['rows'] == duplicate['rows']
    assert response.json()['committed_at'] is None
    assert a.post(url + '/commit').status_code == 409
    assert a.post(url + '/cancel').status_code == 409
    assert a.get(url).json()['status'] == '已取消'
    assert a.get('/api/imports').json()['items'][0]['status'] == '已取消'
    valid = preview(a, [['不应入库', '测试企业', '北京']]).json()
    assert admin.post(f"/api/imports/{valid['id']}/cancel").status_code == 200
    assert a.get('/api/jobs').json()['total'] == 1
    assert any(log['action'] == '取消岗位导入' for log in admin.get('/api/audit-logs').json()['items'])

def test_cancel_and_commit_cannot_both_succeed(clients):
    a, b = clients('alice'), clients('alice')
    batch = preview(a, [['竞争测试', '测试公司', '上海']]).json()
    url = f"/api/imports/{batch['id']}"
    with ThreadPoolExecutor(max_workers=2) as pool:
        commit = pool.submit(a.post, url + '/commit')
        cancel = pool.submit(b.post, url + '/cancel')
        assert sorted([commit.result().status_code, cancel.result().status_code]) == [200, 409]
    status = a.get(url).json()['status']
    assert status in ('已完成', '已取消')
    assert a.get('/api/jobs').json()['total'] == (1 if status == '已完成' else 0)

def test_student_isolation_transfer_resumes_and_history(clients):
    a, b, admin = clients('alice'), clients('bob'), clients('admin')
    student = a.post('/api/students', json={'name': '学员甲', 'major': '计算机'}).json(); sid = student['id']
    assert a.post('/api/students', json={'name': '越权分配', 'owner_id': 3}).status_code == 403
    assert b.get('/api/students').json()['total'] == 0
    assert b.get(f'/api/students/{sid}').status_code == 404
    assert b.put(f'/api/students/{sid}', json={'name': '篡改'}).status_code == 404
    upload = a.post(f'/api/students/{sid}/resumes', files={'file': ('简历.pdf', b'%PDF-1.4\n%%EOF')})
    assert upload.status_code == 201
    rid = upload.json()['id']
    assert b.get(f'/api/resumes/{rid}/download').status_code == 404
    assert b.post(f'/api/students/{sid}/resumes', files={'file': ('简历.pdf', b'%PDF-1.4')}).status_code == 404
    job = add_job(a)
    reco = a.post(f'/api/students/{sid}/recommendations', json={'job_id': job['id']})
    assert reco.status_code == 201
    assert b.post(f'/api/students/{sid}/recommendations', json={'job_id': job['id']}).status_code == 404
    assert a.post(f'/api/students/{sid}/recommendations', json={'job_id': job['id']}).status_code == 409
    assert b.post(f'/api/students/{sid}/followups', json={'content': '越权'}).status_code == 404
    assert a.post(f'/api/students/{sid}/followups', json={'recommendation_id': reco.json()['id'], 'stage': '已投递', 'content': '已确认投递', 'next_date': str(today())}).status_code == 201
    payload = {k: job[k] for k in ['title', 'company', 'city', 'education', 'experience', 'major', 'salary_min', 'salary_max', 'description', 'source', 'expires_on', 'status']}
    payload['title'] = '修改后的岗位'
    assert a.put(f"/api/jobs/{job['id']}", json=payload).status_code == 200
    detail = a.get(f'/api/students/{sid}').json()
    assert detail['recommendations'][0]['snapshot']['title'] == '后端工程师'
    assert detail['recommendations'][0]['stage'] == '已投递'
    assert a.get('/api/dashboard').json()['reminders'][0]['student_id'] == sid
    assert b.get('/api/dashboard').json()['reminders'] == []
    assert admin.put(f'/api/students/{sid}', json={'name': '学员甲', 'owner_id': 3}).status_code == 200
    assert a.get(f'/api/students/{sid}').status_code == 404
    assert a.get(f'/api/resumes/{rid}/download').status_code == 404
    assert b.get(f'/api/students/{sid}').json()['followups'][0]['content'] == '已确认投递'
    assert b.get(f'/api/resumes/{rid}/download').status_code == 200
    assert any('转交学员' in l['action'] for l in admin.get('/api/audit-logs').json()['items'])

def test_cross_student_recommendation_link_rejected(clients):
    a, b = clients('alice'), clients('bob')
    s1 = a.post('/api/students', json={'name': '甲'}).json()['id']
    s2 = b.post('/api/students', json={'name': '乙'}).json()['id']
    j = add_job(a)
    r = b.post(f'/api/students/{s2}/recommendations', json={'job_id': j['id']}).json()['id']
    assert a.post(f'/api/students/{s1}/followups', json={'content': '测试', 'stage': '已投递', 'recommendation_id': r}).status_code == 404

def test_job_filter_expiry_and_unknown_salary(clients):
    c = clients('alice')
    b = preview(c, [['工程师', '甲公司', '上海', '本科', '应届生', '计算机', 8000, 10000], ['设计师', '乙公司', '北京', '大专', '不限', '', None, None, '', '', str(today() - timedelta(days=1))]]).json()
    c.post(f"/api/imports/{b['id']}/commit")
    assert c.get('/api/jobs?status=有效').json()['total'] == 1
    assert c.get('/api/jobs?status=失效').json()['total'] == 1
    assert c.get('/api/jobs?salary_min=9000').json()['total'] == 1
    assert c.get('/api/jobs?salary_min=11000').json()['total'] == 0
    assert c.get('/api/jobs?q=计算机&city=上海&education=本科&experience=应届生').json()['total'] == 1
    assert c.get('/api/jobs?page=0').status_code == 422
    assert c.get('/api/jobs?q=%25').json()['total'] == 0
    student = c.post('/api/students', json={'name': '测试学员'}).json()['id']
    expired = c.get('/api/jobs?status=失效').json()['items'][0]
    assert c.post(f'/api/students/{student}/recommendations', json={'job_id': expired['id']}).status_code == 400

def test_optional_salary_import_update_and_filter(clients):
    c = clients('alice')
    batch = preview(c, [
        ['未知薪资', '测试公司', '上海'],
        ['只有下限', '测试公司', '上海', '本科', '不限', '', 10000, None],
        ['只有上限', '测试公司', '上海', '本科', '不限', '', None, 15000],
        ['完整区间', '测试公司', '上海', '本科', '不限', '', 8000, 12000],
        ['错误区间', '测试公司', '上海', '本科', '不限', '', 15000, 8000],
    ]).json()
    assert batch['counts']['可导入'] == 4
    assert batch['counts']['错误'] == 1
    result = c.post(f"/api/imports/{batch['id']}/commit").json()
    assert result['counts']['已导入'] == 4
    items = {j['title']: j for j in c.get('/api/jobs').json()['items']}
    assert items['未知薪资']['salary_min'] is None and items['未知薪资']['salary_max'] is None
    assert items['只有下限']['salary_min'] == 10000 and items['只有下限']['salary_max'] is None
    assert items['只有上限']['salary_min'] is None and items['只有上限']['salary_max'] == 15000
    assert c.get('/api/jobs?salary_min=9000').json()['total'] == 3
    assert c.get('/api/jobs?salary_min=13000').json()['total'] == 1
    job = items['完整区间']
    payload = {k: job[k] for k in ['title', 'company', 'city', 'education', 'experience', 'major', 'salary_min', 'salary_max', 'description', 'source', 'expires_on', 'status']}
    for lower, upper in [(None, None), (9000, None), (None, 14000)]:
        payload.update(salary_min=lower, salary_max=upper)
        response = c.put(f"/api/jobs/{job['id']}", json=payload)
        assert response.status_code == 200, response.text
        assert response.json()['salary_min'] == lower and response.json()['salary_max'] == upper
    payload.update(salary_min=14000, salary_max=9000)
    assert c.put(f"/api/jobs/{job['id']}", json=payload).status_code == 422
