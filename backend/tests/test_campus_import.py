from datetime import datetime
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

from app.imports import CAMPUS_HEADERS, inspect_xlsx, template_bytes, error_workbook_bytes


def campus_file(rows, headers=CAMPUS_HEADERS, extra_sheets=True):
    wb = Workbook()
    ws = wb.active
    ws.title = '校招汇总表（优先）'
    ws.append(headers)
    for values in rows:
        ws.append([values.get(h) for h in headers])
    if extra_sheets:
        wb.create_sheet('企业名录').append(['公司', '链接'])
        wb.create_sheet('内推码').append(['企业名称', '内推码'])
        wb.active = 1
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def record(**changes):
    return {'公司': '校招公司', '工作地点': '北京, 上海', '招聘岗位': '软件工程师',
            '学历要求': '本科, 硕士, 博士', '截止时间': '招满为止',
            '简历投递链接': 'https://example.com/apply?' + 'x' * 750,
            '招聘届次': '2027届', '备注': '关注原公告', '是否笔试 (1)': '是', **changes}


def test_template_has_exact_campus_headers():
    wb = load_workbook(BytesIO(template_bytes()))
    assert list(next(wb.active.values)) == CAMPUS_HEADERS


def test_campus_sheet_selection_reordered_columns_and_original_fields():
    original = record(**{'招聘岗位': '算法工程师' * 200, '工作地点': '上海, ' * 100, '专业要求': '计算机' * 100})
    row = inspect_xlsx(campus_file([original], list(reversed(CAMPUS_HEADERS))))[0]
    assert row['errors'] == []
    assert row['data']['education'] == '本科'
    assert row['data']['salary_min'] is None
    assert row['data']['experience'] == '不限'
    assert row['data']['expires_on'] is None
    for key, value in original.items():
        assert row['data']['recruitment_details'][key] == value.strip()


@pytest.mark.parametrize('deadline,expected', [('2026/11/10', '2026-11-10'), (datetime(2026, 11, 10), '2026-11-10'), ('尽快投递', None), ('', None)])
def test_deadlines(deadline, expected):
    row = inspect_xlsx(campus_file([record(**{'截止时间': deadline})]))[0]
    assert row['errors'] == []
    assert row['data']['expires_on'] == expected


def test_invalid_rows_and_download_keep_original_columns():
    rows = inspect_xlsx(campus_file([record(**{'截止时间': '2026/4/31'}), record(**{'招聘岗位': ''}), record(**{'备注': '=1+1'})]))
    assert all(r['state'] == '错误' for r in rows)
    wb = load_workbook(BytesIO(error_workbook_bytes(rows)))
    sheet = wb.active
    assert list(next(sheet.values)) == ['原始行号', *CAMPUS_HEADERS, '原因']
    assert sheet['K2'].value == '2026/4/31'
    assert all(c.data_type != 'f' for row in sheet for c in row)


def test_large_campus_batch_and_ambiguous_sheets():
    assert len(inspect_xlsx(campus_file([record()] * 2001))) == 2001
    wb = load_workbook(BytesIO(campus_file([record()])))
    wb.copy_worksheet(wb.worksheets[0])
    output = BytesIO()
    wb.save(output)
    with pytest.raises(HTTPException) as exc:
        inspect_xlsx(output.getvalue())
    assert exc.value.status_code == 400


def test_campus_preview_commit_and_edit_preserve_details(clients):
    client = clients('alice')
    original = record()
    response = client.post('/api/imports/preview', files={'file': ('校招.xlsx', campus_file([original, original]))})
    assert response.status_code == 201, response.text
    batch = response.json()
    assert batch['counts']['可导入'] == batch['counts']['重复'] == 1
    assert client.get('/api/jobs').json()['total'] == 0
    assert client.post(f"/api/imports/{batch['id']}/commit").json()['counts']['已导入'] == 1
    job = client.get('/api/jobs').json()['items'][0]
    assert job['recruitment_details']['简历投递链接'] == original['简历投递链接']
    update = {k: job[k] for k in ('title', 'company', 'city')}
    update['title'] = '更新岗位名称'
    edited = client.put(f"/api/jobs/{job['id']}", json=update)
    assert edited.status_code == 200, edited.text
    assert edited.json()['recruitment_details'] == job['recruitment_details']
