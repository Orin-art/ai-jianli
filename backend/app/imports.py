import hashlib
import re
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from datetime import date, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError
from fastapi import HTTPException
from .schemas import JobInput, Education
from typing import get_args

HEADERS = {'岗位名称': 'title', '公司名称': 'company', '城市': 'city', '学历要求': 'education', '经验要求': 'experience', '专业要求': 'major', '最低月薪（元）': 'salary_min', '最高月薪（元）': 'salary_max', '岗位描述': 'description', '来源': 'source', '截止日期': 'expires_on'}
MAX_ROWS = 2000
CAMPUS_MAX_ROWS = 20000
CAMPUS_HEADERS = ['更新时间', '公司', '企业类型', '行业类别', '批次', '招聘届次', '工作地点', '招聘岗位', '开始时间', '截止时间', '简历投递链接', '公告链接', '是否笔试', '公告来源', '备注', '专业要求', '学历要求', '父记录', '是否笔试 (1)', '父记录 2']
CAMPUS_REQUIRED = {'公司', '工作地点', '招聘岗位'}


def text_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else str(value).strip() if value is not None else ''


def campus_payload(raw):
    details = {key: text_value(value) for key, value in raw.items()}
    data = {target: details.get(source, '') for source, target in [('公司', 'company'), ('工作地点', 'city'), ('招聘岗位', 'title'), ('专业要求', 'major'), ('公告来源', 'source')]}
    data['recruitment_details'] = details
    degrees = [v.strip() for v in re.split(r'[,，、/;；]+', details.get('学历要求', '')) if v.strip()]
    allowed = list(get_args(Education))
    if any(v not in allowed for v in degrees):
        raise ValueError('学历要求：请使用不限、高中及以下、中专、大专、本科、硕士、博士，可用逗号分隔')
    data['education'] = min(degrees, key=allowed.index) if degrees else '不限'
    deadline = raw.get('截止时间')
    if isinstance(deadline, datetime):
        data['expires_on'] = deadline.date()
    elif isinstance(deadline, date):
        data['expires_on'] = deadline
    elif details.get('截止时间') not in (None, '', '/', '招满为止', '尽快投递', '长期有效', '待定'):
        try:
            data['expires_on'] = datetime.strptime(details['截止时间'].replace('/', '-'), '%Y-%m-%d').date()
        except ValueError:
            raise ValueError('截止时间：请填写日期、招满为止、尽快投递、长期有效或待定')
    return data


def error_workbook_bytes(rows):
    campus = any(row['data'].get('recruitment_details') for row in rows)
    headers = CAMPUS_HEADERS if campus else list(HEADERS)
    values = []
    for row in rows:
        if not row['errors']:
            continue
        data = row['data']
        cells = [data.get('recruitment_details', {}).get(h, '') for h in headers] if campus else [data.get(k, '') for k in HEADERS.values()]
        values.append([row['row_number'], *cells, '；'.join(row['errors'])])
    return workbook_bytes(['原始行号', *headers, '原因'], values)

def fingerprint(data):
    values = [str(data[k]).strip().casefold() for k in ('title', 'company', 'city')]
    return hashlib.sha256('\x1f'.join(values).encode()).hexdigest()

def safe_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value

def workbook_bytes(headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = '岗位'
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='257763')
    for row in rows:
        ws.append([safe_cell(x) for x in row])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22
    output = BytesIO()
    wb.save(output)
    return output.getvalue()

def template_bytes():
    wb = load_workbook(BytesIO(workbook_bytes(CAMPUS_HEADERS, [])))
    ws = wb.active
    for column in ('B', 'G', 'H'):
        ws[f'{column}1'].comment = Comment('必填；不要修改表头。相同岗位名称、公司和城市被视为重复。', '知途')
    ws['J1'].comment = Comment('可留空；支持 YYYY-MM-DD、YYYY/MM/DD、Excel 日期、招满为止、尽快投递、长期有效、待定。文字原样保留，不推测日期。', '知途')
    ws['Q1'].comment = Comment('可用逗号分隔多个学历，例如：本科, 硕士, 博士。原文完整保留，筛选按最低学历。', '知途')
    result = BytesIO(); wb.save(result)
    return result.getvalue()

def inspect_xlsx(content, *, with_source=False):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 100 or sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise HTTPException(400, 'Excel 解压后过大或文件结构过于复杂')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError, OSError):
        raise HTTPException(400, '文件不是有效的 .xlsx 工作簿')
    try:
        candidates = []
        for sheet in wb.worksheets:
            sheet.reset_dimensions()
            first_row = next(sheet.iter_rows(), ())
            names = [text_value(c.value) for c in first_row]
            while names and not names[-1]:
                names.pop()
            if names == list(HEADERS) or CAMPUS_REQUIRED.issubset(names):
                candidates.append((sheet, names))
        if len(candidates) != 1:
            raise HTTPException(400, '请保留一个可识别的岗位汇总工作表；表头需包含公司、工作地点、招聘岗位，或使用旧版模板')
        ws, headers = candidates[0]
        campus = CAMPUS_REQUIRED.issubset(headers)
        if len(headers) > 30 or len(set(headers)) != len(headers):
            raise HTTPException(400, '表头重复或列数超过 30')
        if campus and any(h not in CAMPUS_HEADERS for h in headers):
            raise HTTPException(400, '存在不支持的字段，请使用最新模板的列名')
        # Ignore untrusted worksheet dimension metadata; enforce actual streamed limits.
        ws.reset_dimensions()
        iterator = ws.iter_rows()
        first = next(iterator, None)
        if not first or len(first) > 30:
            raise HTTPException(400, '表头为空或列数超过 30')
        result = []
        limit = CAMPUS_MAX_ROWS if campus else MAX_ROWS
        scanned = 0
        for row_number, cells in enumerate(iterator, start=2):
            scanned += 1
            if scanned > limit:
                raise HTTPException(400, f'每次最多导入 {limit} 行，请拆分文件')
            if all(c.value is None for c in cells):
                continue
            if len(cells) > len(headers) and any(c.value is not None for c in cells[len(headers):]):
                raise HTTPException(400, f'第 {row_number} 行存在模板外的数据列')
            raw = {(h if campus else HEADERS[h]): (cells[i].value if i < len(cells) else None) for i, h in enumerate(headers)}
            display = {k: v.isoformat() if isinstance(v, (date, datetime)) else str(v) if v is not None else '' for k, v in raw.items()}
            errors = []
            if any(c.data_type == 'f' for c in cells):
                errors.append('不支持公式，请先粘贴为数值或文本')
            if campus:
                try:
                    data = campus_payload(raw)
                    payload = JobInput.model_validate(data).model_dump(mode='json')
                except ValidationError as exc:
                    names = {'title': '招聘岗位', 'company': '公司', 'city': '工作地点', 'major': '专业要求', 'source': '公告来源'}
                    errors.extend(f"{names.get(str(e['loc'][0]), '岗位')}：{e['msg']}" for e in exc.errors())
                    payload = {'title': display.get('招聘岗位', ''), 'company': display.get('公司', ''), 'city': display.get('工作地点', ''), 'recruitment_details': display}
                except ValueError as exc:
                    errors.append(str(exc))
                    payload = {'title': display.get('招聘岗位', ''), 'company': display.get('公司', ''), 'city': display.get('工作地点', ''), 'recruitment_details': display}
                result.append({'row_number': row_number, 'data': payload, 'errors': errors, 'state': '错误' if errors else '可导入'})
                continue
            data = {}
            for key, value in raw.items():
                if value is None or str(value).strip() == '':
                    continue
                if key == 'expires_on' and isinstance(value, datetime):
                    value = value.date()
                data[key] = value if key in ('salary_min', 'salary_max', 'expires_on') else str(value).strip()
            try:
                job = JobInput.model_validate(data)
                payload = job.model_dump(mode='json')
            except ValidationError as exc:
                names = {v: k for k, v in HEADERS.items()}
                errors.extend(f"{names.get(str(e['loc'][0]), '岗位') if e['loc'] else '岗位'}：{e['msg']}" for e in exc.errors())
                payload = display
            result.append({'row_number': row_number, 'data': payload, 'errors': errors, 'state': '错误' if errors else '可导入'})
        if not result:
            raise HTTPException(400, '工作表没有岗位数据')
        return (result, ws.title) if with_source else result
    finally:
        wb.close()
