"""Export selected source rows without including other jobs or auxiliary sheets."""
from copy import copy, deepcopy
from io import BytesIO
import re

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .imports import CAMPUS_HEADERS, HEADERS, safe_cell, workbook_bytes
from .models import ImportBatch


def copy_rows(source, target, row_numbers):
    headers = [c.value for c in source[1]]
    while headers and headers[-1] is None:
        headers.pop()
    column_count = len(headers)
    mapping = {old: new for new, old in enumerate([1, *sorted(set(row_numbers))], start=1)}
    for attr in ('sheet_format', 'sheet_properties', 'print_options', 'page_margins', 'page_setup', 'HeaderFooter'):
        setattr(target, attr, copy(getattr(source, attr)))
    target.views = deepcopy(source.views)
    for view in target.views.sheetView:
        view.topLeftCell = 'A1'
        for selection in view.selection:
            selection.activeCell = 'A1'
            selection.sqref = 'A1'
    target.sheet_state = 'visible'
    target.freeze_panes = source.freeze_panes
    target.print_title_rows = source.print_title_rows
    target.print_title_cols = source.print_title_cols
    for key, dimension in source.column_dimensions.items():
        target.column_dimensions[key] = copy(dimension)
        target.column_dimensions[key].parent = target
    for old_row, new_row in mapping.items():
        dimension = copy(source.row_dimensions[old_row])
        dimension.index = new_row
        dimension.parent = target
        dimension.hidden = False
        target.row_dimensions[new_row] = dimension
        for old in source.iter_rows(min_row=old_row, max_row=old_row, max_col=column_count):
            for cell in old:
                new = target.cell(new_row, cell.column, safe_cell(cell.value))
                new.font = copy(cell.font)
                new.fill = copy(cell.fill)
                new.border = copy(cell.border)
                new.alignment = copy(cell.alignment)
                new.number_format = cell.number_format
                new.protection = copy(cell.protection)
                if cell.comment:
                    new.comment = copy(cell.comment)
                link = cell.hyperlink.target if cell.hyperlink else None
                if link and re.match(r'^https?://', link, re.I):
                    new.hyperlink = copy(cell.hyperlink)
                elif isinstance(cell.value, str) and re.match(r'^https?://', cell.value, re.I):
                    new.hyperlink = cell.value
    for merged in source.merged_cells.ranges:
        selected = [r for r in mapping if merged.min_row <= r <= merged.max_row]
        if len(selected) == merged.max_row - merged.min_row + 1:
            target.merge_cells(start_row=mapping[merged.min_row], end_row=mapping[merged.max_row],
                               start_column=merged.min_col, end_column=merged.max_col)
    if source.auto_filter.ref:
        target.auto_filter = deepcopy(source.auto_filter)
        target.auto_filter.ref = f'A1:{get_column_letter(column_count)}{len(mapping)}'
        target.auto_filter.sortState = None
        # Previous filters may refer to excluded rows or hide the selected results.
        target.auto_filter.filterColumn = []
    if source.print_area:
        target.print_area = f'A1:{get_column_letter(column_count)}{len(mapping)}'


def add_fallback_sheet(output, jobs, campus):
    headers = CAMPUS_HEADERS if campus else list(HEADERS)
    values = []
    for job in jobs:
        if campus:
            values.append([job.recruitment_details.get(h, '') for h in headers])
        else:
            values.append([getattr(job, key) for key in HEADERS.values()])
    source = load_workbook(BytesIO(workbook_bytes(headers, values)))
    target = output.create_sheet('历史岗位-校招字段' if campus else '历史岗位-旧版字段')
    copy_rows(source.active, target, list(range(2, len(jobs) + 2)))
    source.close()


def student_jobs_workbook(db, jobs):
    output = Workbook()
    output.remove(output.active)
    groups, fallback = {}, {True: [], False: []}
    for job in jobs:
        if job.source_batch_id and job.source_row_number:
            groups.setdefault(job.source_batch_id, []).append(job)
        else:
            fallback[bool(job.recruitment_details)].append(job)
    for batch_id, group in groups.items():
        batch = db.get(ImportBatch, batch_id)
        if not batch or not batch.source_file or not batch.source_sheet:
            for job in group:
                fallback[bool(job.recruitment_details)].append(job)
            continue
        source = load_workbook(BytesIO(batch.source_file), keep_links=False)
        try:
            sheet = source[batch.source_sheet]
            if len(output.worksheets) == 0:
                output.loaded_theme = source.loaded_theme
                output.epoch = source.epoch
            target = output.create_sheet(sheet.title)
            copy_rows(sheet, target, [job.source_row_number for job in group])
        finally:
            source.close()
    for campus, group in fallback.items():
        if group:
            add_fallback_sheet(output, group, campus)
    result = BytesIO()
    output.save(result)
    output.close()
    return result.getvalue()
